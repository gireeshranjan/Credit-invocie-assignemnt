"""Shared paths, validation and Excel exports for the credit-risk assignment."""
from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
from openpyxl.styles import Font, PatternFill

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw"
PROCESSED = ROOT / "data" / "processed"
EXCEL = ROOT / "exports" / "excel"
ANALYTICS = ROOT / "reports" / "analytics"
MODELS = ROOT / "models"
MODEL_REPORTS = ROOT / "reports" / "models"
SEED = 20261026
LIMITATION = ("Synthetic, retrospective default-label classification. Default event dates, "
              "outcome horizon and bureau observation dates are unavailable. Repayment "
              "features may overlap the outcome. Scores are not validated future default probabilities.")

def load_base(path=None):
    path = Path(path) if path else PROCESSED / "analytics_base.csv"
    a = pd.read_csv(path)
    if not a.account_id.notna().all() or not a.account_id.is_unique or a.duplicated().any():
        raise ValueError("Analytics base must have exactly one non-null, unique row per loan.")
    if a.customer_id.isna().any() or not a.target_default.isin([0, 1]).all():
        raise ValueError("Invalid customer keys or target labels.")
    return a

def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, indent=2, ensure_ascii=True, allow_nan=False), encoding="utf-8")

def fingerprint(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def write_workbook(path, tables):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, frame in tables.items():
            if len(name) > 31:
                raise ValueError(f"Excel sheet name too long: {name}")
            frame.to_excel(xl, sheet_name=name, index=False, na_rep="")
            ws = xl.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 30
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
                ws.column_dimensions[cell.column_letter].width = min(42, max(15, len(str(cell.value)) + 2))
            for j, col in enumerate(frame.columns, 1):
                if col.endswith("_rate") or col.endswith("_share") or col.endswith("_probability"):
                    for cells in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                        cells[0].number_format = "0.00%"

def add_segments(a):
    a = a.copy()
    a["income_group"] = pd.cut(a.annual_income,
        [-np.inf, 30000, 60000, 90000, 120000, np.inf], right=False,
        labels=["Under 30k", "30k-under 60k", "60k-under 90k", "90k-under 120k", "120k+"])
    a["age_group"] = pd.cut(a.age, [-np.inf, 30, 40, 50, 60, np.inf], right=False,
                              labels=["Under 30", "30-39", "40-49", "50-59", "60+"])
    for col in ["income_group", "age_group"]:
        a[col] = a[col].cat.add_categories(["Unknown"]).fillna("Unknown")
    return a

def portfolio_summary(a, keys):
    x = a.copy()
    for col in keys:
        if not isinstance(x[col].dtype, pd.CategoricalDtype):
            x[col] = x[col].fillna("Unknown")
    x["_defaulted_amount"] = x.loan_amount.where(x.target_default.eq(1), 0)
    g = x.groupby(keys, observed=True, dropna=False).agg(
        loan_count=("account_id", "size"), customer_count=("customer_id", "nunique"),
        defaults=("target_default", "sum"), default_rate=("target_default", "mean"),
        original_loan_amount=("loan_amount", "sum"),
        original_amount_on_defaulted=("_defaulted_amount", "sum"),
        median_annual_income=("annual_income", "median")).reset_index()
    g["loan_share"] = g.loan_count / len(a)
    g["amount_share"] = g.original_loan_amount / a.loan_amount.sum()
    g["default_share"] = g.defaults / a.target_default.sum() if a.target_default.sum() else np.nan
    g["small_group"] = g.loan_count.lt(100)
    return g

def add_credit_history_groups(a):
    """Prior bureau write-offs and outside-credit activity; no bureau record is its own group."""
    a = a.copy()
    a["write_off_group"] = pd.cut(a.bur_n_written_off, [-1, 0, 1, 2, np.inf], labels=["0", "1", "2", "3+"])
    a["write_off_group"] = a.write_off_group.cat.add_categories(["No bureau records"])
    a.loc[a.bur_thin_file.eq(1), "write_off_group"] = "No bureau records"
    a["credit_activity"] = np.select(
        [a.bur_thin_file.eq(1), a.bur_n_active.eq(0)],
        ["No bureau records", "Bureau history, none active"], "Active outside credit")
    return a

def vintage_trend(a, freq="Q"):
    """Origination cohorts with control limits: overall rate +/- 2 binomial standard errors."""
    x = a.assign(vintage=pd.to_datetime(a.origination_date).dt.to_period(freq).astype(str))
    g = x.groupby("vintage").agg(loan_count=("account_id", "size"), defaults=("target_default", "sum"),
                                 original_loan_amount=("loan_amount", "sum")).reset_index()
    g["default_rate"] = g.defaults / g.loan_count
    overall = a.target_default.mean()
    margin = 2 * np.sqrt(overall * (1 - overall) / g.loan_count)
    g["expected_rate"] = overall
    g["lower_limit"], g["upper_limit"] = (overall - margin).clip(lower=0), overall + margin
    g["outside_limits"] = ~g.default_rate.between(g.lower_limit, g.upper_limit)
    return g

def trend_test(vintages, periods_per_year=4):
    """Least-squares slope of cohort default rates, per year, with its p-value."""
    from scipy.stats import linregress
    fit = linregress(np.arange(len(vintages)), vintages.default_rate)
    return fit.slope * periods_per_year, fit.pvalue

def load_repayment():
    return pd.read_csv(RAW / "repayment_history.csv",
                       usecols=["account_id", "payment_month", "days_past_due", "payment_status"])

def repayment_timeline(a, repayment):
    """Monthly repayments dated as origination month + instalment number (assumed schedule)."""
    r = repayment.merge(a[["account_id", "origination_date", "target_default"]], on="account_id",
                        validate="many_to_one")
    start = pd.to_datetime(r.origination_date)
    month = start.dt.year * 12 + start.dt.month - 1 + r.payment_month
    r["calendar_month"] = (month // 12).astype(str) + "-" + (month % 12 + 1).astype(str).str.zfill(2)
    r["dpd30"] = r.days_past_due.ge(30)
    r["outcome"] = np.where(r.target_default.eq(1), "Defaulted", "Performing")
    return r

def seasoning_curve(r):
    return r.groupby(["payment_month", "outcome"]).agg(
        loan_count=("account_id", "size"), dpd30_rate=("dpd30", "mean"),
        mean_dpd=("days_past_due", "mean")).reset_index()

def calendar_delinquency(r, min_loans=500):
    g = r.groupby("calendar_month").agg(active_loans=("account_id", "size"),
                                        dpd30_rate=("dpd30", "mean")).reset_index()
    g["small_group"] = g.active_loans.lt(min_loans)
    return g
