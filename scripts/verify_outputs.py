"""Verify source preservation, exports, customer isolation and saved-model predictions."""
import json
import numpy as np
import pandas as pd
import joblib
from openpyxl import load_workbook
from portfolio_common import (ROOT, RAW, EXCEL, PROCESSED, MODEL_REPORTS, MODELS,
                              load_base, write_json)
from model_utils import feature_frame, evaluate

def run():
    a = load_base()
    loans = pd.read_csv(RAW / "loan_accounts.csv")
    customers = pd.read_csv(RAW / "customers.csv")
    bureau = pd.read_csv(RAW / "bureau_credit.csv")
    assert len(a) == len(loans) == 14685
    assert set(a.account_id) == set(loans.account_id)
    assert a.customer_id.nunique() == len(customers) == 10000
    joined = a.merge(loans, on="account_id", suffixes=("_clean", "_raw"), validate="one_to_one")
    for col in ["target_default", "loan_amount", "interest_rate_pct", "tenure_months"]:
        np.testing.assert_allclose(joined[col+"_clean"], joined[col+"_raw"], equal_nan=True)
    c = a.drop_duplicates("customer_id").merge(customers, on="customer_id",
            suffixes=("_clean", "_raw"), validate="one_to_one")
    for col in ["num_dependents", "years_at_residence", "annual_income"]:
        assert c[col+"_clean"].isna().equals(c[col+"_raw"].isna())
    originally_missing = c.years_employed_raw.isna()
    logical_zero = originally_missing & c.employment_type_raw.eq("Unemployed")
    assert c.years_employed_clean.isna().equals(originally_missing & ~logical_zero)
    assert c.loc[logical_zero, "years_employed_clean"].eq(0).all()
    expected_bureau = bureau.groupby("customer_id").days_past_due.agg(["size", "count", "max"])
    b = a.drop_duplicates("customer_id").set_index("customer_id")
    assert b.bur_n_lines.eq(expected_bureau["size"].reindex(b.index).fillna(0)).all()
    assert b.bur_n_dpd_known.eq(expected_bureau["count"].reindex(b.index).fillna(0)).all()
    np.testing.assert_allclose(b.bur_max_dpd, expected_bureau["max"].reindex(b.index), equal_nan=True)
    assert b.loc[b.bur_n_dpd_known.eq(0), "bur_max_dpd"].isna().all()
    for name in ["product", "income", "gender", "age", "region", "missed_payments", "bureau_dpd",
                 "write_offs", "credit_activity", "employment", "vintage", "vintage_quarterly"]:
        table = pd.read_csv(ROOT / "reports" / "analytics" / f"{name}.csv")
        assert table.loan_count.sum() == len(a), name
        assert table.defaults.sum() == a.target_default.sum(), name
    split = pd.read_csv(MODEL_REPORTS / "split_membership.csv")
    assert len(split) == len(a) and set(split.account_id) == set(a.account_id)
    assert split.groupby("customer_id").partition.nunique().eq(1).all()
    bysplit = split.set_index("account_id").partition
    validations = pd.read_csv(MODEL_REPORTS / "validation_comparison.csv")
    selection = json.loads((MODEL_REPORTS / "selection_before_test.json").read_text())
    assert selection["test_used_for_selection"] is False
    eligible = validations[validations.auc.ge(validations.auc.max()-.005)].copy()
    eligible["complexity"] = eligible.scenario.eq("enriched").astype(int)*2 + eligible.algorithm.eq("xgboost").astype(int)
    expected = eligible.sort_values(["brier", "complexity", "model"]).iloc[0]["model"]
    bundle = joblib.load(MODELS / "selected_model.joblib")
    assert expected == selection["selected_model"] == bundle["selected_model"]
    features = bundle["feature_spec"]["numeric"] + bundle["feature_spec"]["categorical"]
    assert not set(features) & {"gender", "target_default", "account_id", "customer_id"}
    assert not any(f.startswith("coll_") for f in features)
    scores = pd.read_csv(PROCESSED / "loan_risk_scores.csv")
    assert len(scores) == len(a) and scores.account_id.is_unique
    assert scores.default_probability.between(0, 1).all()
    assert scores.driver_1.notna().all()
    model_p = bundle["pipeline"].predict_proba(feature_frame(a, bundle["feature_spec"]))[:, 1]
    np.testing.assert_allclose(scores.default_probability, model_p, rtol=1e-6, atol=1e-7)
    status_map = {"train": "training_in_sample", "validation": "validation_used_for_selection", "test": "held_out_test"}
    assert scores.evaluation_status.equals(scores.account_id.map(bysplit).map(status_map))
    tests = pd.read_csv(MODEL_REPORTS / "test_comparison.csv").set_index("model")
    predictions = pd.read_csv(MODEL_REPORTS / "test_predictions.csv")
    test_ids = set(split.loc[split.partition.eq("test"), "account_id"])
    for name, frame in predictions.groupby("model"):
        assert set(frame.account_id) == test_ids
        assert frame.account_id.is_unique
        truth = frame.account_id.map(a.set_index("account_id").target_default)
        np.testing.assert_array_equal(frame.target_default, truth)
        metrics = evaluate(frame.target_default, frame.default_probability, tests.loc[name, "threshold"])
        for field in ["auc", "ks", "gini", "precision", "recall", "brier", "accuracy"]:
            np.testing.assert_allclose(metrics[field], tests.loc[name, field], atol=1e-6)
    # Entire cleaned Excel dataset must match CSV, including every blank cell.
    excel = pd.read_excel(EXCEL / "analytics_base.xlsx", sheet_name="analytics_base")
    csv_for_comparison = a.copy()
    for col in ["origination_date", "customer_since"]:
        csv_for_comparison[col] = pd.to_datetime(csv_for_comparison[col])
        excel[col] = pd.to_datetime(excel[col])
    pd.testing.assert_frame_equal(csv_for_comparison, excel, check_dtype=False, check_exact=False, rtol=1e-9, atol=1e-7)
    for name in ["portfolio_analytics.xlsx", "risk_model_results.xlsx", "loan_risk_scores.xlsx"]:
        wb = load_workbook(EXCEL / name, read_only=True, data_only=True)
        assert wb.sheetnames
        for ws in wb:
            assert ws.max_row >= 2 and ws.max_column >= 2
        wb.close()
    refs = json.loads((MODEL_REPORTS / "monitoring_reference.json").read_text())
    assert len(refs) == len(features)+1 and "model_score" in refs
    result = {
        "status": "PASS", "loans_retained": len(a), "customers": int(a.customer_id.nunique()),
        "duplicate_loan_ids": int(a.account_id.duplicated().sum()),
        "customer_cross_partition_leakage": 0, "source_values_and_missingness_preserved": True,
        "excel_matches_cleaned_csv": True, "saved_scores_match_model": True,
        "reported_test_metrics_recomputed": True, "selected_model": expected}
    write_json(ROOT / "reports" / "verification.json", result)
    print(json.dumps(result, indent=2))
    return result

if __name__ == "__main__":
    run()
