"""Rebuild the cleaned loan-level dataset using the agreed notebook policies.
Standalone script: no notebook execution or statistical imputation required.
The two age-based caps are explicit inherited assumptions; raw CSVs are untouched.
"""
from portfolio_common import EXCEL

def display(*args):
    """Notebook-only exploratory displays have no output in batch execution."""
    pass

def run():
    EXCEL.mkdir(parents=True, exist_ok=True)
    from pathlib import Path

    import numpy as np
    import pandas as pd

    pd.set_option("display.max_columns", 50)
    pd.set_option("display.width", 200)

    ROOT = Path(__file__).resolve().parent.parent
    RAW = ROOT / "data" / "raw"
    OUT = ROOT / "data" / "processed"
    REPORTS = ROOT / "reports"
    OUT.mkdir(parents=True, exist_ok=True)
    REPORTS.mkdir(exist_ok=True)

    # Declared reference cutoff for quality checks; not a verified bureau observation date.
    SNAPSHOT_DATE = pd.Timestamp("2023-12-31")

    # Every check is recorded here and exported at the end as the data-quality log.
    dq_log = []

    def log(table, issue, rows, action):
        dq_log.append({"table": table, "issue": issue, "rows_affected": int(rows), "action": action})

    customers   = pd.read_csv(RAW / "customers.csv",          parse_dates=["customer_since"])
    loans       = pd.read_csv(RAW / "loan_accounts.csv",      parse_dates=["origination_date"])
    bureau      = pd.read_csv(RAW / "bureau_credit.csv",      parse_dates=["date_opened", "date_closed"])
    repayment   = pd.read_csv(RAW / "repayment_history.csv")
    collections = pd.read_csv(RAW / "collections_events.csv", parse_dates=["event_date"])

    raw = {"customers": customers, "loans": loans, "bureau": bureau,
           "repayment": repayment, "collections": collections}

    # Trim stray whitespace in all text columns before checking labels.
    for df in raw.values():
        for col in df.columns:
            if pd.api.types.is_string_dtype(df[col]):
                df[col] = df[col].str.strip()

    pd.DataFrame({name: {"rows": len(df), "columns": df.shape[1]} for name, df in raw.items()}).T

    missing = pd.concat(
        {name: df.isna().sum()[df.isna().sum() > 0] for name, df in raw.items()}
    ).rename("n_missing").to_frame()
    missing["pct"] = [f"{n / len(raw[t]):.1%}" for (t, _), n in missing.n_missing.items()]
    missing

    for name, key in [("customers", "customer_id"), ("loans", "account_id"),
                      ("bureau", "bureau_id"), ("collections", "collection_event_id")]:
        duplicate = raw[name][key].duplicated().sum()
        missing_key = raw[name][key].isna().sum()
        log(name, f"duplicate {key}", duplicate, "checked; fail if present")
        log(name, f"missing {key}", missing_key, "checked; fail if present")
        assert duplicate == 0 and missing_key == 0, f"Invalid key in {name}"
    duplicate = repayment.duplicated(["account_id", "payment_month"]).sum()
    log("repayment", "duplicate (account_id, payment_month)", duplicate, "checked; fail if present")
    assert duplicate == 0
    assert repayment[["account_id", "payment_month"]].notna().all().all()
    orphans = {
        "loans -> customers": ~loans.customer_id.isin(customers.customer_id),
        "bureau -> customers": ~bureau.customer_id.isin(customers.customer_id),
        "repayment -> loans": ~repayment.account_id.isin(loans.account_id),
        "collections -> loans": ~collections.account_id.isin(loans.account_id),
    }
    for rel, mask in orphans.items():
        log("keys", f"orphan foreign key ({rel})", mask.sum(), "checked; fail if present")
        assert not mask.any(), f"Unmatched reference: {rel}"
    pd.DataFrame(dq_log)


    print("before:", customers.gender.value_counts().to_dict())
    mask = customers.gender.isin(["M", "F"])
    log("customers", "gender coded 'M'/'F' instead of 'Male'/'Female'", mask.sum(), "mapped M→Male, F→Female")
    customers["gender"] = customers.gender.replace({"M": "Male", "F": "Female"})
    print("after: ", customers.gender.value_counts().to_dict())

    was_missing = customers.years_employed.isna()
    unemp_null = was_missing & customers.employment_type.eq("Unemployed")
    other_null = was_missing & ~unemp_null
    customers["years_employed_missing_original"] = was_missing.astype(int)
    customers["years_employed_imputed"] = unemp_null.astype(int)
    customers.loc[unemp_null, "years_employed"] = 0.0
    customers["years_employed_missing"] = customers.years_employed.isna().astype(int)
    log("customers", "years_employed missing (Unemployed)", unemp_null.sum(),
        "set to 0: current job tenure; flag years_employed_imputed")
    log("customers", "years_employed missing (other employment types)", other_null.sum(),
        "retained as unknown; flag years_employed_missing; no statistical imputation")
    print("Logical zero fills:", int(unemp_null.sum()))
    print("Remaining unknown tenure:", int(customers.years_employed.isna().sum()))


    max_years = (customers.age - 18).clip(lower=0)
    too_long = customers.years_employed > max_years
    display(customers.loc[too_long, ["customer_id", "age", "years_employed", "employment_type"]].head())

    customers["years_employed_capped"] = too_long.astype(int)
    customers.loc[too_long, "years_employed"] = max_years[too_long]
    log("customers", "years_employed exceeds age − 18", too_long.sum(), "capped at age − 18 + flag years_employed_capped")
    print(f"capped {too_long.sum():,} rows")

    for col in ["num_dependents", "years_at_residence"]:
        n = customers[col].isna()
        customers[f"{col}_missing"] = n.astype(int)
        log("customers", f"{col} missing", n.sum(), f"retained as unknown + flag {col}_missing")
    too_long = customers.years_at_residence > customers.age
    customers["years_at_residence_capped"] = too_long.astype(int)
    customers.loc[too_long, "years_at_residence"] = customers.age[too_long]
    log("customers", "years_at_residence exceeds age", too_long.sum(),
        "capped at age + flag years_at_residence_capped")
    customers[["years_employed", "num_dependents", "years_at_residence"]].isna().sum()


    floor = customers.annual_income == customers.annual_income.min()
    log("customers", f"annual_income at floor value {customers.annual_income.min():,.0f} (looks clipped)",
        floor.sum(), "kept as-is; noted in assumptions log")

    customers["age_band"] = pd.cut(customers.age, [0, 29, 39, 49, 59, 200],
                                   labels=["21-29", "30-39", "40-49", "50-59", "60+"])
    customers["income_band"] = pd.cut(customers.annual_income, [-np.inf, 30_000, 60_000, 90_000, 120_000, np.inf], right=False,
                                      labels=["<30k", "30-60k", "60-90k", "90-120k", "120k+"])
    customers[["age_band", "income_band"]].apply(lambda s: s.value_counts().sort_index()).fillna(0).astype(int)

    n = loans.interest_rate_pct.isna()
    loans["interest_rate_missing"] = n.astype(int)
    log("loans", "interest_rate_pct missing", n.sum(),
        "retained as unknown + flag interest_rate_missing")
    loans.groupby("loan_type").interest_rate_pct.agg(["count", "median"])


    expected = loans.loan_type.isin(["Mortgage", "Auto Loan"]).astype(int)
    log("loans", "collateral_flag inconsistent with loan_type", (loans.collateral_flag != expected).sum(),
        "checked (rule: Mortgage/Auto = secured)")

    since = loans.customer_id.map(customers.set_index("customer_id").customer_since)
    bad = loans.origination_date < since
    loans["dq_orig_before_customer_since"] = bad.astype(int)
    log("loans", "origination_date earlier than customer_since", bad.sum(),
        "flagged; customer_since treated as unreliable, not used as a feature")
    print(f"loans originated before the customer relationship 'began': {bad.sum():,} of {len(loans):,} ({bad.mean():.0%})")

    loans["vintage_year"] = loans.origination_date.dt.year

    bureau["dpd_missing"] = bureau.days_past_due.isna().astype(int)
    for status, group in bureau.groupby("account_status"):
        log("bureau", f"days_past_due missing on {status} lines", group.days_past_due.isna().sum(),
            "retained as unknown; no DPD imputation; tracked in aggregate missing/known counts")
    log("bureau", "date_closed blank on Active lines",
        (bureau.account_status.eq("Active") & bureau.date_closed.isna()).sum(),
        "retained as not applicable while account is active")
    print("DPD values kept unknown:", int(bureau.days_past_due.isna().sum()))


    fut_open = bureau.date_opened > SNAPSHOT_DATE
    fut_close = bureau.date_closed > SNAPSHOT_DATE
    print(f"opened after {SNAPSHOT_DATE.date()}: {fut_open.sum():,} ({fut_open.mean():.0%})")
    print(f"closed after {SNAPSHOT_DATE.date()}: {fut_close.sum():,} — latest {bureau.date_closed.max().date()}")
    log("bureau", f"date_opened after snapshot {SNAPSHOT_DATE.date()}", fut_open.sum(),
        "kept for descriptive analytics; observation timing unverified; no assumed application-time availability")
    log("bureau", "date_closed in the future (up to 2029)", fut_close.sum(), "kept; same as above")

    status_vs_close = (((bureau.account_status == "Active") & bureau.date_closed.notna())
                       | ((bureau.account_status != "Active") & bureau.date_closed.isna()))
    log("bureau", "account_status inconsistent with date_closed", status_vs_close.sum(), "checked")
    log("bureau", "outstanding_amount > sanctioned_amount",
        (bureau.outstanding_amount > bureau.sanctioned_amount).sum(), "checked")

    display(repayment.groupby("payment_status").agg(
        rows=("account_id", "size"),
        dpd_min=("days_past_due", "min"), dpd_max=("days_past_due", "max"),
        share_paid_zero=("emi_paid", lambda s: (s == 0).mean())))

    rules = {
        "On-Time with days_past_due > 0": (repayment.payment_status == "On-Time") & (repayment.days_past_due > 0),
        "Missed with emi_paid > 0":       (repayment.payment_status == "Missed") & (repayment.emi_paid > 0),
        "emi_paid > emi_due":             repayment.emi_paid > repayment.emi_due + 0.01,
        "negative amounts":               (repayment[["emi_due", "emi_paid", "outstanding_balance"]] < 0).any(axis=1),
    }
    for issue, mask in rules.items():
        log("repayment", issue, mask.sum(), "checked")
        print(f"{issue:<32} {mask.sum()}")

    months = repayment.groupby("account_id").payment_month.nunique()
    log("repayment", "accounts without exactly 12 months", (months != 12).sum(), "checked")
    repayment = repayment.sort_values(["account_id", "payment_month"]).reset_index(drop=True)

    orig = collections.account_id.map(loans.set_index("account_id").origination_date)
    bad = collections.event_date < orig
    collections["dq_event_before_origination"] = bad.astype(int)
    log("collections", "event_date earlier than loan origination_date", bad.sum(),
        "flagged; collections used for analytics only, not as model features")
    print(f"events dated before the loan existed: {bad.sum():,} of {len(collections):,} ({bad.mean():.0%})")

    has_coll = loans.account_id.isin(collections.account_id)
    print(f"default rate — accounts with collections: {loans[has_coll].target_default.mean():.1%}, "
          f"without: {loans[~has_coll].target_default.mean():.1%}")

    b = bureau.assign(
        is_active=bureau.account_status.eq("Active").astype(int),
        is_written_off=bureau.account_status.eq("Written-Off").astype(int),
        # Counts of observed 30+ DPD; missing DPD is tracked separately.
        is_dpd30=bureau.days_past_due.ge(30).astype(int),
        is_nonbank=bureau.lender_category.isin(["NBFC", "Fintech"]).astype(int),
        active_sanctioned=bureau.sanctioned_amount.where(bureau.account_status.eq("Active"), 0),
        active_outstanding=bureau.outstanding_amount.where(bureau.account_status.eq("Active"), 0),
    )
    bureau_agg = b.groupby("customer_id").agg(
        bur_n_lines=("bureau_id", "size"),
        bur_n_active=("is_active", "sum"),
        bur_n_written_off=("is_written_off", "sum"),
        bur_total_sanctioned=("sanctioned_amount", lambda s: s.sum(min_count=1)),
        bur_total_outstanding=("outstanding_amount", lambda s: s.sum(min_count=1)),
        bur_active_sanctioned=("active_sanctioned", lambda s: s.sum(min_count=1)),
        bur_active_outstanding=("active_outstanding", lambda s: s.sum(min_count=1)),
        bur_max_dpd=("days_past_due", "max"),
        bur_mean_dpd=("days_past_due", "mean"),
        bur_n_dpd30=("is_dpd30", "sum"),
        bur_nonbank_share=("is_nonbank", "mean"),
        bur_n_dpd_known=("days_past_due", "count"),
        bur_n_dpd_missing=("dpd_missing", "sum"),
    )
    bureau_agg["bur_has_written_off"] = bureau_agg.bur_n_written_off.gt(0).astype(int)
    bureau_agg["bur_dpd_coverage"] = bureau_agg.bur_n_dpd_known / bureau_agg.bur_n_lines
    bureau_agg["bur_utilisation"] = (
        bureau_agg.bur_active_outstanding / bureau_agg.bur_active_sanctioned.where(
            bureau_agg.bur_active_sanctioned > 0))
    bureau_agg = bureau_agg.drop(columns=["bur_active_sanctioned", "bur_active_outstanding"])
    print(bureau_agg.shape)
    bureau_agg.head()


    r = repayment.assign(
        missed=(repayment.payment_status == "Missed").astype(int),
        late=(repayment.payment_status == "Late").astype(int),
    )
    # Longest run of consecutive missed months: number each run, take its length where missed = 1.
    new_run = (r.account_id != r.account_id.shift()) | (r.missed != r.missed.shift())
    run_len = r.groupby(new_run.cumsum()).missed.transform("size") * r.missed

    repay_agg = r.groupby("account_id").agg(
        rp_months=("payment_month", "size"),
        rp_emi=("emi_due", "median"),
        rp_n_missed=("missed", "sum"),
        rp_n_late=("late", "sum"),
        rp_max_dpd=("days_past_due", "max"),
        rp_mean_dpd=("days_past_due", "mean"),
        rp_total_due=("emi_due", "sum"),
        rp_total_paid=("emi_paid", "sum"),
    )
    repay_agg["rp_max_consec_missed"] = run_len.groupby(r.account_id).max()
    repay_agg["rp_pct_on_time"] = 1 - (repay_agg.rp_n_missed + repay_agg.rp_n_late) / repay_agg.rp_months
    repay_agg["rp_paid_to_due"] = repay_agg.rp_total_paid / repay_agg.rp_total_due
    repay_agg["rp_dpd_first3"] = r[r.payment_month <= 3].groupby("account_id").days_past_due.mean()
    repay_agg["rp_dpd_last3"] = r[r.payment_month >= 10].groupby("account_id").days_past_due.mean()
    repay_agg["rp_dpd_trend"] = repay_agg.rp_dpd_last3 - repay_agg.rp_dpd_first3
    repay_agg["rp_ever_30dpd"] = (repay_agg.rp_max_dpd >= 30).astype(int)
    repay_agg["rp_last_outstanding"] = r[r.payment_month == 12].set_index("account_id").outstanding_balance
    repay_agg = repay_agg.drop(columns=["rp_total_due", "rp_total_paid"])
    print(repay_agg.shape)
    repay_agg.head()

    c = collections.assign(
        legal=(collections.action_type == "Legal Notice").astype(int),
        wo_referral=(collections.action_type == "Written-Off Referral").astype(int),
    )
    coll_agg = c.groupby("account_id").agg(
        coll_n_events=("collection_event_id", "size"),
        coll_total_recovered=("amount_recovered", lambda s: s.sum(min_count=1)),
        coll_has_legal=("legal", "max"),
        coll_has_wo_referral=("wo_referral", "max"),
        coll_n_recovered_missing=("amount_recovered", lambda s: s.isna().sum()),
        coll_n_event_before_origination=("dq_event_before_origination", "sum"),
    )
    print(coll_agg.shape)
    coll_agg.head()

    base = (loans
            .merge(customers, on="customer_id", how="left", validate="many_to_one")
            .merge(bureau_agg, on="customer_id", how="left", validate="many_to_one")
            .merge(repay_agg, on="account_id", how="left", validate="one_to_one")
            .merge(coll_agg, on="account_id", how="left", validate="one_to_one"))
    no_bureau = base.bur_n_lines.isna()
    base["bur_thin_file"] = no_bureau.astype(int)
    bureau_count_cols = [
        "bur_n_lines", "bur_n_active", "bur_n_written_off", "bur_n_dpd30",
        "bur_n_dpd_known", "bur_n_dpd_missing", "bur_has_written_off",
    ]
    base.loc[no_bureau, bureau_count_cols] = 0
    base[bureau_count_cols] = base[bureau_count_cols].astype(int)
    log("join", "accounts whose customer has no bureau record", no_bureau.sum(),
        "recorded counts 0 + bur_thin_file; DPD, amounts, shares and ratios remain unknown")
    no_collections = base.coll_n_events.isna()
    base["coll_no_records"] = no_collections.astype(int)
    # Zero for an absent group; preserve unknown amounts within an existing group.
    base.loc[no_collections, coll_agg.columns] = 0
    log("join", "accounts with no collections records", no_collections.sum(),
        "zero recorded events/recovery + coll_no_records; no claim about unrecorded activity")
    base["loan_to_income"] = base.loan_amount / base.annual_income.where(base.annual_income > 0)
    base["emi_to_income"] = base.rp_emi * 12 / base.annual_income.where(base.annual_income > 0)
    base["customer_n_loans"] = base.groupby("customer_id").account_id.transform("size")
    assert len(base) == len(loans), "join changed the row count"
    assert base.account_id.notna().all() and base.account_id.is_unique
    assert not base.duplicated().any(), "complete duplicate rows found"
    assert set(base.account_id) == set(loans.account_id)
    print(f"rows: {len(loans):,} in -> {len(base):,} out | columns: {base.shape[1]}")
    print("Complete duplicate rows:", int(base.duplicated().sum()))


    print(f"shape: {base.shape}")
    print(f"unknown cells retained: {int(base.isna().sum().sum())}")
    print(f"default rate: {base.target_default.mean():.1%}")
    print(f"original loan amounts total: {base.loan_amount.sum():,.0f}")
    coverage = pd.DataFrame({
        "column": base.columns,
        "known_rows": base.notna().sum().to_numpy(),
        "missing_rows": base.isna().sum().to_numpy(),
    })
    coverage["known_share"] = coverage.known_rows / len(base)
    display(coverage.loc[coverage.missing_rows.gt(0)])
    for col in ["num_dependents", "years_at_residence"]:
        assert customers[col].isna().sum() == int(missing.loc[("customers", col), "n_missing"])
    assert loans.interest_rate_pct.isna().sum() == int(missing.loc[("loans", "interest_rate_pct"), "n_missing"])
    assert bureau.days_past_due.isna().sum() == int(missing.loc[("bureau", "days_past_due"), "n_missing"])
    assert customers.years_employed.isna().sum() == (
        int(missing.loc[("customers", "years_employed"), "n_missing"])
        - int(customers.years_employed_imputed.sum()))
    assert base.loc[base.bur_n_dpd_known.eq(0), ["bur_max_dpd", "bur_mean_dpd"]].isna().all().all()
    assert (base.bur_n_dpd_known + base.bur_n_dpd_missing).eq(base.bur_n_lines).all()
    # Demonstrate inclusion of Unknown customers in analytics.
    tenure_group = pd.cut(base.years_employed, [-1, 0, 2, 5, 10, np.inf],
                          labels=["0", ">0-2", ">2-5", ">5-10", ">10"])
    tenure_group = tenure_group.cat.add_categories(["Unknown"]).fillna("Unknown")
    tenure_summary = base.groupby(tenure_group, observed=True).agg(
        loan_count=("account_id", "size"),
        original_loan_amount=("loan_amount", "sum"),
        recorded_default_rate=("target_default", "mean"),
    )
    assert tenure_summary.loan_count.sum() == len(base)
    display(tenure_summary)


    dq = pd.DataFrame(dq_log)
    dq

    source_of = {
        **{c: "loan_accounts" for c in loans.columns},
        **{c: "customers" for c in customers.columns if c != "customer_id"},
        **{c: "bureau_credit (aggregated per customer)" for c in list(bureau_agg.columns) + ["bur_thin_file"]},
        **{c: "repayment_history (aggregated per account)" for c in repay_agg.columns},
        **{c: "collections_events (aggregated, analytics only)" for c in coll_agg.columns},
    }
    derived = {
        "years_employed_imputed": "1 if missing tenure for an Unemployed customer was filled with 0",
        "years_employed_capped": "1 if years_employed was capped at age − 18",
        "num_dependents_missing": "1 if number of dependents remains unknown",
        "years_employed_missing_original": "1 if tenure was missing in the source",
        "years_employed_missing": "1 if employment tenure remains unknown",
        "years_at_residence_capped": "1 if residence tenure was capped at age",
        "years_at_residence_missing": "1 if residence tenure remains unknown",
        "interest_rate_missing": "1 if contractual interest rate remains unknown",
        "dq_orig_before_customer_since": "1 if origination_date < customer_since (data issue)",
        "age_band": "age bucket", "income_band": "annual income bucket", "vintage_year": "origination year",
        "bur_n_lines": "number of external credit lines", "bur_n_active": "active external lines",
        "bur_n_written_off": "written-off external lines", "bur_has_written_off": "1 if any recorded bureau line is Written-Off; 0 means none recorded; consult bur_thin_file",
        "bur_total_sanctioned": "sum of recorded external approved amounts; blank if no bureau records", "bur_total_outstanding": "sum of recorded external outstanding balances; blank if no bureau records",
        "bur_utilisation": "active outstanding / active sanctioned; blank if denominator is zero or unavailable", "bur_max_dpd": "maximum observed bureau DPD; blank if no DPD is known",
        "bur_mean_dpd": "mean observed bureau DPD; consult coverage; blank if none known", "bur_n_dpd30": "count of observed bureau lines at 30+ DPD; unknown DPD excluded",
        "bur_nonbank_share": "share of lines from NBFC/Fintech lenders", "bur_n_dpd_missing": "recorded bureau lines with unknown DPD",
        "bur_n_dpd_known": "bureau lines with observed DPD",
        "bur_dpd_coverage": "observed DPD count / bureau line count; blank if no bureau lines",
        "bur_thin_file": "1 if customer has no bureau record",
        "rp_months": "months of history (always 12)", "rp_emi": "monthly instalment",
        "rp_n_missed": "missed payments in 12m", "rp_n_late": "late payments in 12m",
        "rp_max_dpd": "worst DPD in 12m", "rp_mean_dpd": "average DPD in 12m",
        "rp_max_consec_missed": "longest run of consecutive missed months",
        "rp_pct_on_time": "share of months paid on time", "rp_paid_to_due": "total paid / total due",
        "rp_dpd_first3": "average DPD months 1–3", "rp_dpd_last3": "average DPD months 10–12",
        "rp_dpd_trend": "last3 − first3 DPD (positive = worsening)", "rp_ever_30dpd": "1 if ever 30+ DPD",
        "rp_last_outstanding": "balance after month 12",
        "coll_n_events": "number of recorded collections actions",
        "coll_no_records": "1 if no collections records match this loan",
        "coll_n_recovered_missing": "collections records with unknown recovered amount",
        "coll_n_event_before_origination": "collections events dated before loan origination", "coll_total_recovered": "amount recovered",
        "coll_has_legal": "1 if legal notice issued", "coll_has_wo_referral": "1 if referred for write-off",
        "loan_to_income": "loan_amount / annual_income", "emi_to_income": "annual EMI / annual_income",
        "customer_n_loans": "number of loans this customer holds",
    }
    raw_descriptions = {}
    for line in (RAW / "DATA_DICTIONARY.md").read_text(encoding="utf-8").splitlines():
        parts = [part.strip() for part in line.split("|")]
        if len(parts) >= 4 and parts[1]:
            raw_descriptions[parts[1].strip("*")] = parts[2].strip("*")
    dictionary = pd.DataFrame({
        "column": base.columns,
        "source": [source_of.get(c, "derived") for c in base.columns],
        "description": [derived.get(c, raw_descriptions.get(c, "See source dictionary")) for c in base.columns],
        "dtype": base.dtypes.astype(str).values,
    })


    policy = pd.DataFrame([
        ("years_employed", "Missing + Unemployed", "0 + logical-fill flag", "Current employment tenure."),
        ("years_employed", "Missing + other employment types", "Keep null + flag", "Unknown group; retain all loans."),
        ("num_dependents", "Missing", "Keep null + flag", "Unknown is not zero dependents."),
        ("years_at_residence", "Missing", "Keep null + flag", "Unknown is not a recent move."),
        ("interest_rate_pct", "Missing", "Keep null + flag", "Unknown is not an interest-free loan."),
        ("bureau.days_past_due", "Missing in any status", "Keep null; observed-only summaries + coverage",
         "No invented delinquency. All-missing summaries stay null."),
        ("bureau.date_closed", "Blank while Active", "Keep null / not applicable", "Closure has not occurred."),
        ("bureau summaries", "No matching records", "Zero recorded counts; unknown amounts/DPD/ratios",
         "Unavailable external history is not proof of no external debt."),
        ("collections summaries", "No matching records", "Zero recorded events/recovery + absence flag",
         "Unknown amounts in existing records remain unknown."),
        ("model preprocessing", "Features requiring imputation", "Fit any imputer on training data only",
         "No model-specific statistical fills in the analytics dataset."),
    ], columns=["field", "condition", "treatment", "reason"])
    base.to_csv(OUT / "analytics_base.csv", index=False, na_rep="")
    dq.to_csv(REPORTS / "data_quality_log.csv", index=False)
    coverage.to_csv(REPORTS / "analytics_missing_coverage.csv", index=False)
    dictionary.to_csv(REPORTS / "analytics_column_dictionary.csv", index=False)
    policy.to_csv(REPORTS / "missing_value_policy.csv", index=False)
    with pd.ExcelWriter(EXCEL / "analytics_base.xlsx", engine="openpyxl") as xl:
        for name, frame in {
            "analytics_base": base, "data_quality_log": dq, "column_dictionary": dictionary,
            "missing_coverage": coverage, "missing_value_policy": policy,
        }.items():
            frame.to_excel(xl, sheet_name=name, index=False, na_rep="")
            ws = xl.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                ws.column_dimensions[cell.column_letter].width = min(48, max(16, len(str(cell.value)) + 3))
        from openpyxl.styles import Font, PatternFill, Alignment
        for ws in xl.book.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.alignment = Alignment(wrap_text=True)
            ws.row_dimensions[1].height = 32
    import importlib.util
    if any(importlib.util.find_spec(module) for module in ["pyarrow", "fastparquet"]):
        base.to_parquet(OUT / "analytics_base.parquet", index=False)
        print("Parquet written.")
    else:
        print("Parquet skipped: optional engine unavailable. Use analytics_base.csv.")
    for filename in ["analytics_base.csv", "analytics_base.xlsx"]:
        print(f"Wrote {(EXCEL if filename.endswith('.xlsx') else OUT) / filename}")

    return base

if __name__ == "__main__":
    run()
